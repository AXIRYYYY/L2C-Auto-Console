"""银行流水按电子日记账标签拆分。

用途：按电子日记账 A 列标签（通常为凭证号）把银行流水 PDF 逐笔拆分为独立文件，
      并在行数与页数明显不匹配时拒绝执行（防止文件错配）。

输入：脚本同目录下「恰好一个」银行流水 PDF + 「恰好一个」含「日记账」的 xlsx
输出：与原 PDF 同级的「分割完成/{标签}.pdf」

用法：python bankStatementSplitterBasedOnJournal.py（无交互，自动扫描脚本所在目录）

依赖：pypdf, openpyxl
对应回放：docs/demo/replay/bank_split.json
"""

import os
import sys

from pypdf import PdfReader, PdfWriter
import openpyxl
from openpyxl import load_workbook

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")


def check_unique_pdf():
    """
    检查是否有唯一的 PDF 文件。
    """
    script_dir = os.path.dirname(os.path.abspath(__file__))  # 脚本所在目录（= 数据目录）
    pdf_files = [file for file in os.listdir(script_dir) if file.endswith('.pdf')]
    if len(pdf_files)!= 1:
        print("未找到唯一的 PDF 文件，请确保有且仅有一个 PDF 文件与脚本放在同一文件夹下！")
        return False
    return True


def check_journal_xlsx():
    """
    检查是否有唯一包含'日记账'的 Excel 文件。
    """
    script_dir = os.path.dirname(os.path.abspath(__file__))  # 脚本所在目录（= 数据目录）
    xlsx_files = [file for file in os.listdir(script_dir) if '日记账' in file and file.endswith('.xlsx')]
    if len(xlsx_files)!= 1:
        print("未找到唯一包含'日记账'的 Excel 文件，请确保有且仅有一个相应文件与脚本放在同一文件夹下！")
        return False
    return True


def get_pdf_and_xlsx_paths():
    """
    获取唯一的 PDF 文件和包含'日记账'的 Excel 文件的路径。
    """
    script_dir = os.path.dirname(os.path.abspath(__file__))  # 脚本所在目录（= 数据目录）
    pdf_files = [file for file in os.listdir(script_dir) if file.endswith('.pdf')]
    xlsx_files = [file for file in os.listdir(script_dir) if '日记账' in file and file.endswith('.xlsx')]
    return os.path.join(script_dir, pdf_files[0]), os.path.join(script_dir, xlsx_files[0])


def split_pdfs_by_tag(input_pdf_path, xlsx_path):
    """
    根据 Excel 文件第一列的标记拆分 PDF 文件，每个标记对应一个新的 PDF 文件，并保存到指定文件夹。
    """
    workbook = load_workbook(xlsx_path)
    sheet = workbook.active
    start_row = 5
    end_row = sheet.max_row - 2
    pdf_reader = PdfReader(open(input_pdf_path, 'rb'))
    num_pages = len(pdf_reader.pages)
    if abs(end_row - start_row + 1 - num_pages) > 2:
        print("Excel 文件行数与 PDF 页面数相差过大，请检查文件数据！")
        return
    tag_pages_dict = {}
    for row in range(start_row, end_row + 1):
        tag = sheet.cell(row=row, column=1).value
        if tag not in tag_pages_dict:
            tag_pages_dict[tag] = []
        tag_pages_dict[tag].append(row - start_row)

    output_folder = os.path.join(os.path.dirname(input_pdf_path), "分割完成")
    if not os.path.exists(output_folder):
        os.mkdir(output_folder)

    for tag, page_indices in tag_pages_dict.items():
        pdf_writer = PdfWriter()
        for index in page_indices:
            pdf_writer.add_page(pdf_reader.pages[index])
        output_path = os.path.join(output_folder, f"{tag}.pdf")
        with open(output_path, 'wb') as outfile:
            pdf_writer.write(outfile)


if __name__ == "__main__":
    if not check_unique_pdf() or not check_journal_xlsx():
        exit(1)
    pdf_file, xlsx_file = get_pdf_and_xlsx_paths()
    split_pdfs_by_tag(pdf_file, xlsx_file)
    print("处理完成，生成的文件已保存到 '分割完成' 文件夹中。")