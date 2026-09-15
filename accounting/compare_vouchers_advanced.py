"""新旧凭证清单深度比对（Jaccard 内容签名）。

用途：比对新旧两版凭证清单，输出 内容修改 / 仅凭证号修改 / 编号内容均修改 / 新增 / 删除 五类差异，
      并对「附单据数」变化行做红色高亮（审计核对用）。

输入：两份 Excel（修改前 / 修改后），列结构见 EXPECTED_COLUMNS
输出：凭证比对结果_{时间戳}/ 目录（含最终报告 xlsx 与源文件副本）

用法：
    python compare_vouchers_advanced.py                          # 弹窗选择两份文件
    python compare_vouchers_advanced.py 修改前.xlsx 修改后.xlsx   # 命令行直接运行（无弹窗）

依赖：pandas, openpyxl
对应回放：docs/demo/replay/voucher_compare.json
"""

import pandas as pd
import tkinter as tk
from tkinter import filedialog, messagebox
import os
import shutil
import sys
import subprocess
from typing import List, Dict, Tuple, Any, Set
from datetime import datetime

# =============================================================================
# 0. 全局配置 (Global Configuration)
# =============================================================================
SIMILARITY_THRESHOLD = 0.5
DISSIMILARITY_THRESHOLD = 0.2

CLI_FILES = [a for a in sys.argv[1:] if a.strip()]
USE_CLI = len(CLI_FILES) >= 2


def notify(title: str, message: str, error: bool = False) -> None:
    """命令行模式下打印日志，GUI 模式下弹窗"""
    if USE_CLI:
        print(f"[{title}] {message}")
        return
    (messagebox.showerror if error else messagebox.showinfo)(title, message)

# =============================================================================
# 1. 核心功能函数 (Core Functions)
# =============================================================================

def select_file(title: str) -> str:
    """打开文件选择对话框 (不再单独创建root，复用主程序的root)"""
    # 注意：这里去掉了 root = tk.Tk()，改由 main 函数统一管理
    filepath = filedialog.askopenfilename(title=title, filetypes=[("Excel files", "*.xlsx *.xls")])
    if not filepath: 
        messagebox.showerror("错误", "未选择文件，程序即将退出。")
        sys.exit() # 使用 sys.exit() 确保彻底退出
    return filepath

def load_and_clean_data(filepath: str, columns: List[str]) -> pd.DataFrame:
    """加载Excel文件，进行数据清洗和预处理"""
    try:
        df = pd.read_excel(filepath, engine='openpyxl')
    except Exception as e:
        notify("文件读取错误", f"无法读取文件: {filepath}\n错误: {e}", error=True); sys.exit()
    for col in columns:
        if col not in df.columns: df[col] = None
    amount_cols = ['借方金额', '贷方金额', '数量', '单价', '外币金额', '汇率']
    for col in amount_cols:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors='coerce').fillna(0)
            if col not in ['数量', '汇率']: df[col] = df[col].round(2)
    str_cols = [col for col in df.columns if col not in amount_cols]
    for col in str_cols:
        df[col] = df[col].fillna('').astype(str).str.strip()
    if '凭证日期' in df.columns:
        df['凭证日期'] = pd.to_datetime(df['凭证日期'], errors='coerce').dt.strftime('%Y-%m-%d').fillna('')
    df['凭证ID'] = df['凭证类别'].astype(str) + '-' + df['凭证号'].astype(str) + '-' + df['凭证日期'].astype(str)
    return df

def get_voucher_content_map(df: pd.DataFrame, content_cols: List[str]) -> Dict[str, frozenset]:
    """生成从 凭证ID 到 其内容签名(frozenset) 的映射"""
    temp_df = df.copy()
    temp_df['行内容签名'] = temp_df[content_cols].apply(tuple, axis=1)
    grouped_signatures = temp_df.groupby('凭证ID')['行内容签名'].apply(frozenset)
    return grouped_signatures.to_dict()

def create_content_to_vouchers_map(voucher_to_content_map: Dict[str, frozenset]) -> Dict[frozenset, List[str]]:
    """创建反向的 "内容 -> [凭证ID列表]" 的映射"""
    content_to_vouchers = {}
    for voucher_id, content_hash in voucher_to_content_map.items():
        if not content_hash: continue
        if content_hash not in content_to_vouchers:
            content_to_vouchers[content_hash] = []
        content_to_vouchers[content_hash].append(voucher_id)
    return content_to_vouchers

def calculate_jaccard_similarity(set1: Set, set2: Set) -> float:
    """计算两个集合的Jaccard相似度"""
    if not set1 and not set2: return 1.0
    intersection = len(set1.intersection(set2))
    union = len(set1.union(set2))
    return intersection / union if union != 0 else 0.0

# =============================================================================
# 2. 核心比对与报告函数
# =============================================================================

def compare_data(df_before: pd.DataFrame, df_after: pd.DataFrame, content_cols: List[str]) -> Dict[str, Any]:
    """核心比对逻辑"""
    summary_cols = ['凭证ID', '凭证类别', '凭证号', '凭证日期']
    before_lookup = df_before[summary_cols].drop_duplicates(subset=['凭证ID']).set_index('凭证ID')
    after_lookup = df_after[summary_cols].drop_duplicates(subset=['凭证ID']).set_index('凭证ID')
    
    v_to_c_before = get_voucher_content_map(df_before, content_cols)
    v_to_c_after = get_voucher_content_map(df_after, content_cols)

    ids_before = set(v_to_c_before.keys())
    ids_after = set(v_to_c_after.keys())
    common_ids = ids_before & ids_after
    
    initial_content_modified_ids = [vid for vid in common_ids if v_to_c_before.get(vid) != v_to_c_after.get(vid)]
    true_content_modified_ids, reclassified_deleted, reclassified_added = [], [], []
    
    for vid in initial_content_modified_ids:
        similarity = calculate_jaccard_similarity(v_to_c_before.get(vid, set()), v_to_c_after.get(vid, set()))
        if similarity < DISSIMILARITY_THRESHOLD:
            reclassified_deleted.append(vid); reclassified_added.append(vid)
        else:
            true_content_modified_ids.append(vid)

    potential_deleted_ids = (ids_before - ids_after) | set(reclassified_deleted)
    potential_added_ids = (ids_after - ids_before) | set(reclassified_added)
    
    number_modified_pairs = []
    processed_deleted, processed_added = set(), set()

    c_to_v_after = create_content_to_vouchers_map(v_to_c_after)

    for vid_del in potential_deleted_ids:
        content_hash = v_to_c_before.get(vid_del)
        if content_hash and content_hash in c_to_v_after:
            for vid_add_candidate in c_to_v_after[content_hash]:
                if vid_add_candidate in potential_added_ids and vid_add_candidate not in processed_added:
                    before_info = before_lookup.loc[vid_del]
                    after_info = after_lookup.loc[vid_add_candidate]
                    number_modified_pairs.append({
                        '凭证类别_修改前': before_info['凭证类别'], '凭证号_修改前': before_info['凭证号'], '凭证日期_修改前': before_info['凭证日期'],
                        '凭证类别_修改后': after_info['凭证类别'], '凭证号_修改后': after_info['凭证号'], '凭证日期_修改后': after_info['凭证日期'],
                        '凭证ID_修改前': vid_del, '凭证ID_修改后': vid_add_candidate
                    })
                    processed_deleted.add(vid_del); processed_added.add(vid_add_candidate)
                    break

    unmatched_deleted_ids = list(potential_deleted_ids - processed_deleted)
    unmatched_added_ids = list(potential_added_ids - processed_added)
    id_and_content_modified_pairs = []
    
    if unmatched_deleted_ids and unmatched_added_ids:
        print(f"正在对 {len(unmatched_deleted_ids)} 个已删除和 {len(unmatched_added_ids)} 个已新增的凭证进行相似度匹配...")
        scores = []
        for del_id in unmatched_deleted_ids:
            for add_id in unmatched_added_ids:
                score = calculate_jaccard_similarity(v_to_c_before.get(del_id, set()), v_to_c_after.get(add_id, set()))
                if score >= SIMILARITY_THRESHOLD: scores.append((score, del_id, add_id))
        
        scores.sort(key=lambda x: x[0], reverse=True)
        for score, del_id, add_id in scores:
            if del_id not in processed_deleted and add_id not in processed_added:
                before_info, after_info = before_lookup.loc[del_id], after_lookup.loc[add_id]
                id_and_content_modified_pairs.append({
                    '凭证类别_修改前': before_info['凭证类别'], '凭证号_修改前': before_info['凭证号'], '凭证日期_修改前': before_info['凭证日期'],
                    '凭证类别_修改后': after_info['凭证类别'], '凭证号_修改后': after_info['凭证号'], '凭证日期_修改后': after_info['凭证日期'],
                    '相似度': round(score, 4), '凭证ID_修改前': del_id, '凭证ID_修改后': add_id
                })
                processed_deleted.add(del_id); processed_added.add(add_id)

    final_deleted_ids = list(potential_deleted_ids - processed_deleted)
    final_added_ids = list(potential_added_ids - processed_added)

    return {
        "content_modified": true_content_modified_ids,
        "number_modified": pd.DataFrame(number_modified_pairs),
        "id_and_content_modified": pd.DataFrame(id_and_content_modified_pairs),
        "added": final_added_ids,
        "deleted": final_deleted_ids
    }

def sort_results_df(df: pd.DataFrame, date_col: str, num_col: str) -> pd.DataFrame:
    if df.empty: return df
    df_copy = df.copy()
    df_copy['凭证号_numeric'] = pd.to_numeric(df_copy[num_col], errors='coerce')
    df_copy.sort_values(by=[date_col, '凭证号_numeric'], na_position='last', inplace=True)
    return df_copy.drop(columns=['凭证号_numeric']).reset_index(drop=True)

def create_comparison_view(writer: pd.ExcelWriter, sheet_name: str, pairs_df: pd.DataFrame, df_before: pd.DataFrame, df_after: pd.DataFrame):
    if pairs_df.empty: return
    all_modified, original_cols = [], [col for col in df_before.columns if col != '凭证ID']
    for _, row in pairs_df.iterrows():
        vid_before, vid_after = row['凭证ID_修改前'], row['凭证ID_修改后']
        before_v = df_before[df_before['凭证ID'] == vid_before][original_cols].reset_index(drop=True)
        after_v = df_after[df_after['凭证ID'] == vid_after][original_cols].reset_index(drop=True)
        comparison_df = pd.merge(before_v, after_v, how='outer', left_index=True, right_index=True, suffixes=('_修改前', '_修改后')).fillna('')
        title = f"修改前ID: {vid_before}  |  修改后ID: {vid_after}"
        if '相似度' in row and pd.notna(row['相似度']): title += f"  |  内容相似度: {row['相似度']:.2%}"
        header_df = pd.DataFrame([title], columns=[original_cols[0]+'_修改前'])
        all_modified.extend([header_df, comparison_df, pd.DataFrame([['']])])
    if all_modified:
        pd.concat(all_modified, ignore_index=True).to_excel(writer, sheet_name=sheet_name, index=False)

def finalize_report_formatting(output_path: str):
    try:
        from openpyxl import load_workbook
        from openpyxl.styles import Font
        workbook = load_workbook(output_path)
        red_font = Font(color="FF0000")
        target_sheets = ['编号内容均修改(详细)', '仅内容修改(详细)', '仅凭证号修改(详细)']
        for sheet_name in workbook.sheetnames:
            ws = workbook[sheet_name]
            for col in ws.columns:
                max_length = max((len(str(cell.value)) for cell in col if cell.value), default=0)
                ws.column_dimensions[col[0].column_letter].width = min(max_length + 2, 50)
            if sheet_name in target_sheets:
                header = [cell.value for cell in ws[1]]
                try:
                    before_col_idx = header.index('附单据数_修改前')
                    after_col_idx = header.index('附单据数_修改后')
                except ValueError: continue
                for row in ws.iter_rows(min_row=2):
                    cell_before, cell_after = row[before_col_idx], row[after_col_idx]
                    val_before = str(cell_before.value) if cell_before.value is not None else ""
                    val_after = str(cell_after.value) if cell_after.value is not None else ""
                    if val_before != val_after and (cell_before.value is not None or cell_after.value is not None):
                        if '修改前ID:' not in str(row[0].value or ''):
                            cell_before.font = red_font; cell_after.font = red_font
        workbook.save(output_path)
    except Exception as e: print(f"\n在最终格式化报告时发生错误: {e}")

def generate_final_report(output_path: str, df_before: pd.DataFrame, df_after: pd.DataFrame, results: Dict[str, Any]):
    with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
        summary_data = { "比对项目": ["内容修改的凭证数", "仅凭证号修改的凭证数", "编号和内容均修改的凭证数", "新增的凭证数", "删除的凭证数"],
                         "数量": [len(results["content_modified"]), len(results["number_modified"]), len(results["id_and_content_modified"]), len(results["added"]), len(results["deleted"])]}
        pd.DataFrame(summary_data).to_excel(writer, sheet_name="比对结果摘要", index=False)
        sorted_id_content_mod = sort_results_df(results["id_and_content_modified"], '凭证日期_修改后', '凭证号_修改后')
        sorted_num_mod = sort_results_df(results["number_modified"], '凭证日期_修改后', '凭证号_修改后')
        content_mod_df = pd.DataFrame({'凭证ID': results["content_modified"]})
        if not content_mod_df.empty:
            summary_info = df_after[['凭证ID', '凭证类别', '凭证号', '凭证日期']].drop_duplicates('凭证ID')
            content_mod_df = pd.merge(content_mod_df, summary_info, on='凭证ID', how='left')
            sorted_content_mod = sort_results_df(content_mod_df, '凭证日期', '凭证号')
        else: sorted_content_mod = pd.DataFrame()

        if not sorted_id_content_mod.empty:
              create_comparison_view(writer, '编号内容均修改(详细)', sorted_id_content_mod, df_before, df_after)
              sorted_id_content_mod.drop(columns=['凭证ID_修改前', '凭证ID_修改后']).to_excel(writer, sheet_name='编号内容均修改(汇总)', index=False)
        if not sorted_content_mod.empty:
            df_to_pass = sorted_content_mod.copy(); df_to_pass['凭证ID_修改前'] = df_to_pass['凭证ID_修改后'] = df_to_pass['凭证ID']
            create_comparison_view(writer, '仅内容修改(详细)', df_to_pass, df_before, df_after)
            sorted_content_mod[['凭证类别', '凭证号', '凭证日期']].to_excel(writer, sheet_name='仅内容修改(汇总)', index=False)
        if not sorted_num_mod.empty:
            create_comparison_view(writer, '仅凭证号修改(详细)', sorted_num_mod, df_before, df_after)
            sorted_num_mod.drop(columns=['凭证ID_修改前', '凭证ID_修改后']).to_excel(writer, sheet_name='仅凭证号修改(汇总)', index=False)
        if results["added"]:
            added_df = df_after[df_after['凭证ID'].isin(results["added"])].copy()
            added_df = sort_results_df(added_df, '凭证日期', '凭证号')
            added_df.drop(columns=['凭证ID']).to_excel(writer, sheet_name='新增的凭证', index=False)
        if results["deleted"]:
            deleted_df = df_before[df_before['凭证ID'].isin(results["deleted"])].copy()
            deleted_df = sort_results_df(deleted_df, '凭证日期', '凭证号')
            deleted_df.drop(columns=['凭证ID']).to_excel(writer, sheet_name='删除的凭证', index=False)

# =============================================================================
# 3. 主程序入口 (Main Execution Block)
# =============================================================================
def main():
    """主执行函数"""
    
    root = None
    if not USE_CLI:
        root = tk.Tk()
        root.withdraw()
    
    try:
        EXPECTED_COLUMNS = ['凭证类别', '凭证号', '凭证日期', '附单据数', '摘要', '科目编码', '科目名称', '借方金额', '贷方金额', '项目编码', '项目', '客户编码', '客户', '供应商编码', '供应商', '部门编码', '部门', '员工编码', '员工', '存货编码', '存货', '规格型号', '数量', '计量单位', '单价', '外币金额', '币种', '汇率', '制单人', '审核人']
        CONTENT_COLUMNS = [col for col in EXPECTED_COLUMNS if col not in ['凭证类别', '凭证号', '凭证日期', '制单人', '审核人', '附单据数']]

        if USE_CLI:
            file_before, file_after = CLI_FILES[0], CLI_FILES[1]
            print(f"[CLI] 修改前: {file_before}\n[CLI] 修改后: {file_after}")
        else:
            file_before = select_file("请选择【修改前】的凭证清单Excel文件")
            file_after = select_file("请选择【修改后】的凭证清单Excel文件")

        print("\n正在加载和清理数据，请稍候...")
        df_before = load_and_clean_data(file_before, EXPECTED_COLUMNS)
        df_after = load_and_clean_data(file_after, EXPECTED_COLUMNS)
        print("数据加载完成。")

        print("正在进行深度数据比对，此过程可能需要一些时间...")
        comparison_results = compare_data(df_before, df_after, CONTENT_COLUMNS)
        print("比对完成。")

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_base_dir = os.path.dirname(file_after)
        result_folder_name = f"凭证比对结果_{timestamp}"
        result_folder_path = os.path.abspath(os.path.join(output_base_dir, result_folder_name))

        os.makedirs(result_folder_path, exist_ok=True)
        print(f"\n已创建结果文件夹: {result_folder_path}")

        output_filename = f"凭证比对分析报告(最终版)_{timestamp}.xlsx"
        output_path = os.path.join(result_folder_path, output_filename)
        
        print(f"正在生成报告到: {output_path}")
        generate_final_report(output_path, df_before, df_after, comparison_results)
        
        print("正在对报告进行最终格式化（调整列宽和高亮差异）...")
        finalize_report_formatting(output_path)
        
        try:
            shutil.copy2(file_before, result_folder_path)
            print(f"已将【修改前】文件复制至结果文件夹。")
            shutil.copy2(file_after, result_folder_path)
            print(f"已将【修改后】文件复制至结果文件夹。")
        except Exception as e:
            notify("警告", f"复制源文件时发生错误: {e}\n\n结果报告已生成，但源文件未能自动归档。", error=True)

        notify("完成", f"比对完成！\n\n所有相关文件已归档至新文件夹:\n{result_folder_path}")
        print("\n比对报告生成完毕！所有文件已归档。")
        
        if not USE_CLI:
            print(f"操作完成，正在尝试打开结果文件夹: {result_folder_path}")
            try:
                if sys.platform == "win32":
                    os.startfile(result_folder_path)
                elif sys.platform == "darwin": # macOS
                    subprocess.run(["open", result_folder_path], check=True)
                else: # Linux
                    subprocess.run(["xdg-open", result_folder_path], check=True)
            except Exception as e:
                print(f"无法自动打开文件夹，请手动访问。\n错误: {e}")

    finally:
        if root is not None:
            root.destroy()

if __name__ == "__main__":
    main()