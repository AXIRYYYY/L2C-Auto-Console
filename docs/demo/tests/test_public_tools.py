"""公开层可验证测试。

覆盖三部分：
  1. 敏感扫描器规则（硬违规能抓到 / 白名单不误报）
  2. 公开文件清单与白名单一致性（public_files.txt ↔ public.gitignore）
  3. 4 个公开脚本的无头集成（CLI 模式真实执行样例数据并断言产物）

运行：python -m pytest docs/demo/tests -q
"""

import io
import os
import shutil
import subprocess
import sys
from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parents[3]
DEMO = REPO / "docs" / "demo"
SAMPLE = REPO / "sample_data"
PY = sys.executable

sys.path.insert(0, str(DEMO))
import sensitive_check as sc  # noqa: E402


def scan_hard(text):
    """返回该文本命中的硬违规类别列表"""
    labels = []
    for label, pat in sc.HARD_PATTERNS:
        if pat.search(text):
            labels.append(label)
    return labels


# ---------------------------------------------------------------- 扫描器规则

# 泄露样例在运行时拼接：本文件自身也在公开清单内，不能包含真实样式的字面量
def _s(*parts):
    return "".join(parts)


LEAK_SAMPLES = [
    (_s("K", "2CLBX202603010001"), "真实主体/内部代号黑名单"),
    (_s("k", "2clbx202603010001"), "真实主体/内部代号黑名单"),
    (_s("路径 M:", "\\", "coding", "\\", "proj", "\\", "a.py"), "本机工作区路径"),
    (_s("C:", "\\", "Users", "\\", "demo", "\\", "secret.txt"), "本机工作区路径"),
    (_s("联系 yej", "min@fox", "mail.com"), "邮箱地址"),
    (_s("电话 138", "0013", "8000"), "手机号"),
    (_s("内部文档 DEMO", "_GUIDE 提到"), "内部意图词"),
    (_s("pass", 'word = "abc"'), "密钥/口令赋值"),
]


@pytest.mark.parametrize("sample_text, expect_label", LEAK_SAMPLES)
def test_hard_patterns_catch_leaks(sample_text, expect_label):
    assert expect_label in scan_hard(sample_text)


@pytest.mark.parametrize(
    "safe_text",
    [
        "SHA256 去重",
        r"D:\your_data\column_info.xlsx",
        r"D:\test\tt\pdf",
        r"C:\Windows\Fonts\simhei.ttf",
        "KZCLBX202603010001",
        "样本编号 202603010001",
    ],
)
def test_safe_tokens_not_flagged_as_hard(safe_text):
    assert scan_hard(safe_text) == []


# -------------------------------------------------------- 清单 / 白名单一致性

def test_manifest_entries_exist():
    missing = []
    for rel in sc.read_manifest():
        source = sc.resolve_source(rel)
        if not (REPO / source).exists():
            missing.append(rel)
    assert not missing, f"清单登记但文件不存在: {missing}"


def test_gitignore_in_sync_with_manifest():
    expected = sc.build_gitignore_text(sc.read_manifest())
    actual = io.open(sc.public_ignore_path(), encoding="utf-8").read()
    assert actual == expected, "public.gitignore 与清单不一致，请运行 sensitive_check.py --sync"


def test_full_scan_passes():
    proc = subprocess.run([PY, str(DEMO / "sensitive_check.py")], capture_output=True)
    assert proc.returncode == 0, proc.stdout.decode("utf-8", "replace")[-1500:]


# ------------------------------------------------------------ 脚本无头集成测试

pytestmark_sample = pytest.mark.skipif(not SAMPLE.exists(), reason="需要 sample_data（先运行 generate_sample_data.py）")


def run_cli(script_rel, args, cwd):
    env = dict(os.environ)
    env["PYTHONIOENCODING"] = "utf-8"
    proc = subprocess.run(
        [PY, str(cwd / script_rel), *args],
        cwd=str(cwd), capture_output=True, text=True, env=env,
        encoding="utf-8", errors="replace", timeout=180,
    )
    return proc


def stage(tmp_path, script_rel):
    dst = tmp_path / script_rel
    dst.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(REPO / script_rel, dst)
    return dst


@pytestmark_sample
def test_compare_vouchers_cli(tmp_path):
    script = "accounting/compare_vouchers_advanced.py"
    stage(tmp_path, script)
    indir = tmp_path / "in"
    indir.mkdir()
    before = SAMPLE / "归档" / "凭证清单_修改前.xlsx"
    after = SAMPLE / "归档" / "凭证清单_修改后.xlsx"
    shutil.copy2(before, indir)
    shutil.copy2(after, indir)

    proc = run_cli(script, [str(indir / before.name), str(indir / after.name)], tmp_path)
    assert proc.returncode == 0, proc.stdout[-1500:]

    reports = list(indir.glob("凭证比对结果_*/凭证比对分析报告*.xlsx"))
    assert reports, "未生成比对报告"

    from openpyxl import load_workbook

    wb = load_workbook(reports[0])
    ws = wb["比对结果摘要"]
    summary = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)}
    assert summary["内容修改的凭证数"] == 1
    assert summary["仅凭证号修改的凭证数"] == 1
    assert summary["新增的凭证数"] == 2
    assert summary["删除的凭证数"] == 1


def _generator_font():
    import importlib.util

    spec = importlib.util.spec_from_file_location("gen_sample", SAMPLE / "generate_sample_data.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod.FONT_FILE


@pytestmark_sample
def test_split_vouchers_cli(tmp_path):
    if _generator_font() is None:
        pytest.skip("环境缺少 CJK 字体，样例 PDF 的中文文本不可被 pypdf 提取")
    script = "accounting/split_vouchers_凭证PDF拆分.py"
    stage(tmp_path, script)
    indir = tmp_path / "in"
    indir.mkdir()
    voucher = SAMPLE / "归档" / "记账凭证-202601.pdf"
    shutil.copy2(voucher, indir)

    proc = run_cli(script, [str(indir / voucher.name)], tmp_path)
    assert proc.returncode == 0, proc.stdout[-1500:]

    splits = sorted((indir / "记账凭证-202601_拆分结果").glob("*.pdf"))
    assert len(splits) == 2
    from pypdf import PdfReader

    pages = sorted(len(PdfReader(str(p)).pages) for p in splits)
    assert pages == [1, 2]


@pytestmark_sample
def test_bank_splitter(tmp_path):
    script = "accounting/bankStatementSplitterBasedOnJournal.py"
    stage(tmp_path, script)
    shutil.copy2(SAMPLE / "归档" / "银行流水-202601.pdf", tmp_path)
    shutil.copy2(SAMPLE / "归档" / "日记账202601.xlsx", tmp_path)

    proc = run_cli(script, [], tmp_path)
    assert proc.returncode == 0, proc.stdout[-1500:]
    outs = sorted((tmp_path / "分割完成").glob("*.pdf"))
    assert len(outs) == 3


@pytestmark_sample
def test_namelistget_cli(tmp_path):
    script = "integration/namelistget_feishu_v9.py"
    stage(tmp_path, script)
    shutil.copy2(SAMPLE / "月底" / "namelist.xlsx", tmp_path)
    shutil.copytree(SAMPLE / "月底" / "附件", tmp_path / "src")
    shutil.copytree(SAMPLE / "月底" / "申请单", tmp_path / "app")

    proc = run_cli(script, [str(tmp_path / "app"), str(tmp_path / "src")], tmp_path)
    assert proc.returncode == 0, proc.stdout[-1500:]

    result = tmp_path / "搜索结果"
    assert (result / "张三-住宿费.pdf").exists(), "未生成合并资料包"
    assert list(result.glob("关键词与PDF文件对应表.xlsx")), "未生成对账表"
